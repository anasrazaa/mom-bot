"""Document export module.

Generates DOCX and PDF versions of a MoMDocument.
"""
import io
from datetime import datetime
from pathlib import Path
from typing import Tuple
from loguru import logger

from app.models.schemas import MoMDocument
from app.config import settings


# ─────────────────────────────────────────────────────────────────────────────
# DOCX
# ─────────────────────────────────────────────────────────────────────────────

def _safe_table_style(table, style: str):
    """Set table style, silently ignoring missing styles in custom templates."""
    try:
        table.style = style
    except (KeyError, Exception):
        pass  # Style not present in template; use default


def _safe_paragraph(doc, text: str, style: str = None):
    """Add a paragraph with an optional style, falling back to default if missing."""
    if style:
        try:
            return doc.add_paragraph(text, style=style)
        except (KeyError, Exception):
            pass
    return doc.add_paragraph(text)


def generate_docx(mom: MoMDocument) -> bytes:
    """Return DOCX bytes for the given MoM."""
    try:
        return _generate_docx_inner(mom)
    except Exception as exc:
        logger.error(f"DOCX generation failed: {exc}", exc_info=True)
        raise RuntimeError(f"DOCX generation failed: {exc}") from exc


def _generate_docx_inner(mom: MoMDocument) -> bytes:
    from docx import Document
    from docx.shared import Inches

    doc, using_template = _create_docx_document_with_template(Document)
    try:
        _apply_template_dynamic_fields(doc, mom)
    except Exception as exc:
        logger.warning(f"Template field replacement failed: {exc}")

    # ── Page margins ────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    # ── Title block ─────────────────────────────────────────────────────────
    if using_template:
        doc.add_paragraph()
    else:
        _add_centered_para(doc, "MINUTES OF MEETING", bold=True, size=16)
        _add_hr(doc)

    # ── Meeting metadata table ───────────────────────────────────────────────
    table = doc.add_table(rows=5, cols=2)
    _safe_table_style(table, "Table Grid")
    _set_cell(table, 0, 0, "Meeting Title", bold=True)
    _set_cell(table, 0, 1, mom.meeting_title)
    _set_cell(table, 1, 0, "Date", bold=True)
    _set_cell(table, 1, 1, mom.date)
    _set_cell(table, 2, 0, "Time", bold=True)
    _set_cell(table, 2, 1, mom.time)
    _set_cell(table, 3, 0, "Venue", bold=True)
    _set_cell(table, 3, 1, mom.venue)
    _set_cell(table, 4, 0, "Chaired By", bold=True)
    _set_cell(table, 4, 1, mom.chaired_by)
    doc.add_paragraph()

    # ── Attendees ────────────────────────────────────────────────────────────
    if mom.attendees:
        _heading(doc, "1. ATTENDEES")
        for name in mom.attendees:
            _safe_paragraph(doc, name, style="List Bullet")
        doc.add_paragraph()

    # ── Agenda ────────────────────────────────────────────────────────────────
    if mom.agenda_items:
        _heading(doc, "2. AGENDA")
        for i, item in enumerate(mom.agenda_items, 1):
            doc.add_paragraph(f"{i}. {item}")
        doc.add_paragraph()

    # ── Discussion Summary ────────────────────────────────────────────────────
    if mom.discussion_summary:
        _heading(doc, "3. DISCUSSION SUMMARY")
        for point in mom.discussion_summary:
            p = doc.add_paragraph()
            run = p.add_run(f"{point.topic}: ")
            run.bold = True
            p.add_run(point.summary)
            if point.speaker:
                p.add_run(f"  [Led by: {point.speaker}]").italic = True
        doc.add_paragraph()

    # ── Decisions ─────────────────────────────────────────────────────────────
    if mom.decisions:
        _heading(doc, "4. DECISIONS TAKEN")
        for i, d in enumerate(mom.decisions, 1):
            text = f"{i}. {d.decision}"
            if d.made_by:
                text += f"  (Approved by: {d.made_by})"
            doc.add_paragraph(text)
        doc.add_paragraph()

    # ── Action Items ─────────────────────────────────────────────────────────
    if mom.action_items:
        _heading(doc, "5. ACTION ITEMS")
        ai_table = doc.add_table(rows=1 + len(mom.action_items), cols=3)
        _safe_table_style(ai_table, "Table Grid")
        headers = ["Action Item", "Responsible Person", "Deadline"]
        for col, h in enumerate(headers):
            _set_cell(ai_table, 0, col, h, bold=True)
        for row_idx, ai in enumerate(mom.action_items, 1):
            _set_cell(ai_table, row_idx, 0, ai.item)
            _set_cell(ai_table, row_idx, 1, ai.responsible)
            _set_cell(ai_table, row_idx, 2, ai.deadline)
        doc.add_paragraph()

    # ── Next Meeting ──────────────────────────────────────────────────────────
    _heading(doc, "6. NEXT MEETING")
    doc.add_paragraph(mom.next_meeting or "To be announced")
    doc.add_paragraph()

    # ── Closing Remarks ───────────────────────────────────────────────────────
    if mom.closing_remarks:
        _heading(doc, "7. CLOSING REMARKS")
        doc.add_paragraph(mom.closing_remarks)
        doc.add_paragraph()

    # ── Additional Notes ──────────────────────────────────────────────────────
    if mom.additional_notes:
        _heading(doc, "8. ADDITIONAL NOTES")
        doc.add_paragraph(mom.additional_notes)
        doc.add_paragraph()

    # ── Footer / Signatures ───────────────────────────────────────────────────
    _add_hr(doc)
    doc.add_paragraph()
    sig_table = doc.add_table(rows=3, cols=2)
    _safe_table_style(sig_table, "Table Grid")
    _set_cell(sig_table, 0, 0, "Prepared By", bold=True)
    _set_cell(sig_table, 0, 1, "Approved By", bold=True)
    _set_cell(sig_table, 1, 0, "_______________________", )
    _set_cell(sig_table, 1, 1, "_______________________")
    _set_cell(sig_table, 2, 0, f"Secretary  |  {mom.date}")
    _set_cell(sig_table, 2, 1, f"{mom.chaired_by}  |  Chairperson")

    doc.add_paragraph()
    gen_p = doc.add_paragraph(
        f"Document generated on: {mom.generated_at.strftime('%d %B %Y, %I:%M %p')}",
    )
    if gen_p.runs:
        gen_p.runs[0].italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# PDF
# ─────────────────────────────────────────────────────────────────────────────

def generate_pdf(mom: MoMDocument) -> bytes:
    """Return PDF bytes for the given MoM."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=1.25 * inch,
        rightMargin=1.25 * inch,
        topMargin=1 * inch,
        bottomMargin=1 * inch,
    )

    styles = getSampleStyleSheet()
    story = []

    # Styles
    title_style = ParagraphStyle("Title2", parent=styles["Title"], fontSize=18, spaceAfter=4)
    inst_style = ParagraphStyle("Inst", parent=styles["Normal"], fontSize=11, alignment=1, spaceAfter=2)
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=12, spaceAfter=6, spaceBefore=12)
    body = styles["Normal"]
    bold_body = ParagraphStyle("BoldBody", parent=body, fontName="Helvetica-Bold")

    story.append(Paragraph("GHULAM ISHAQ KHAN INSTITUTE", inst_style))
    story.append(Paragraph("OF ENGINEERING SCIENCES AND TECHNOLOGY", inst_style))
    story.append(Spacer(1, 8))
    story.append(Paragraph("MINUTES OF MEETING", title_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.black))
    story.append(Spacer(1, 10))

    # Metadata table
    meta_data = [
        ["Meeting Title", mom.meeting_title],
        ["Date", mom.date],
        ["Time", mom.time],
        ["Venue", mom.venue],
        ["Chaired By", mom.chaired_by],
    ]
    meta_table = Table(meta_data, colWidths=[1.5 * inch, 4.5 * inch])
    meta_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.lightgrey),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.whitesmoke, colors.white]),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 12))

    # Attendees
    if mom.attendees:
        story.append(Paragraph("1. ATTENDEES", h1))
        for name in mom.attendees:
            story.append(Paragraph(f"• {name}", body))
        story.append(Spacer(1, 8))

    # Agenda
    if mom.agenda_items:
        story.append(Paragraph("2. AGENDA", h1))
        for i, item in enumerate(mom.agenda_items, 1):
            story.append(Paragraph(f"{i}. {item}", body))
        story.append(Spacer(1, 8))

    # Discussion
    if mom.discussion_summary:
        story.append(Paragraph("3. DISCUSSION SUMMARY", h1))
        for d in mom.discussion_summary:
            story.append(Paragraph(f"<b>{d.topic}:</b> {d.summary}", body))
            if d.speaker:
                story.append(Paragraph(f"<i>Led by: {d.speaker}</i>", body))
        story.append(Spacer(1, 8))

    # Decisions
    if mom.decisions:
        story.append(Paragraph("4. DECISIONS TAKEN", h1))
        for i, d in enumerate(mom.decisions, 1):
            text = f"{i}. {d.decision}"
            if d.made_by:
                text += f" <i>(Approved by: {d.made_by})</i>"
            story.append(Paragraph(text, body))
        story.append(Spacer(1, 8))

    # Action items
    if mom.action_items:
        story.append(Paragraph("5. ACTION ITEMS", h1))
        ai_data = [["#", "Action Item", "Responsible", "Deadline"]]
        for i, ai in enumerate(mom.action_items, 1):
            ai_data.append([str(i), ai.item, ai.responsible, ai.deadline])
        ai_table = Table(ai_data, colWidths=[0.3 * inch, 2.8 * inch, 1.5 * inch, 1.2 * inch])
        ai_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c3e50")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightyellow]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(ai_table)
        story.append(Spacer(1, 8))

    # Next meeting
    story.append(Paragraph("6. NEXT MEETING", h1))
    story.append(Paragraph(mom.next_meeting or "To be announced", body))
    story.append(Spacer(1, 8))

    if mom.closing_remarks:
        story.append(Paragraph("7. CLOSING REMARKS", h1))
        story.append(Paragraph(mom.closing_remarks, body))
        story.append(Spacer(1, 8))

    if mom.additional_notes:
        story.append(Paragraph("8. ADDITIONAL NOTES", h1))
        story.append(Paragraph(mom.additional_notes, body))
        story.append(Spacer(1, 8))

    # Footer
    story.append(HRFlowable(width="100%", thickness=1, color=colors.grey))
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        f"Document generated on {mom.generated_at.strftime('%d %B %Y, %I:%M %p')}",
        ParagraphStyle("footer", parent=body, fontSize=8, textColor=colors.grey, alignment=2),
    ))

    doc.build(story)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# Save helpers
# ─────────────────────────────────────────────────────────────────────────────

def save_docx(mom: MoMDocument) -> Path:
    path = settings.EXPORTS_DIR / f"{mom.meeting_id}_MoM.docx"
    path.write_bytes(generate_docx(mom))
    logger.info(f"DOCX saved: {path}")
    return path


def save_pdf(mom: MoMDocument) -> Path:
    path = settings.EXPORTS_DIR / f"{mom.meeting_id}_MoM.pdf"
    path.write_bytes(generate_pdf(mom))
    logger.info(f"PDF saved: {path}")
    return path


# ─────────────────────────────────────────────────────────────────────────────
# python-docx helpers
# ─────────────────────────────────────────────────────────────────────────────

def _add_centered_para(doc, text: str, bold=False, size=12):
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p


def _heading(doc, text: str):
    from docx.shared import Pt, RGBColor
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    return p


def _set_cell(table, row: int, col: int, text: str, bold: bool = False):
    cell = table.cell(row, col)
    cell.text = text
    if bold:
        for run in cell.paragraphs[0].runs:
            run.bold = True


def _add_hr(doc):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def _create_docx_document_with_template(document_cls):
    """Create a DOCX document, preferring the official template when available.

    The template can contain institutional header/footer/logo and body blocks.
    Body content is preserved as-is so official header layouts remain intact.
    """
    template_path = settings.MOM_DOCX_TEMPLATE_PATH
    if template_path and Path(template_path).exists():
        try:
            doc = document_cls(str(template_path))
            logger.info(f"Using MoM DOCX template: {template_path}")
            return doc, True
        except Exception as exc:
            logger.warning(f"Failed to load MoM DOCX template ({template_path}): {exc}. Falling back to default layout.")
    return document_cls(), False


def _apply_template_dynamic_fields(doc, mom: MoMDocument):
    """Replace supported placeholders in template headers/footers/body.

    Supported placeholders:
      {{MEETING_TITLE}}, {{DATE}}, {{TIME}}, {{VENUE}}, {{CHAIRED_BY}}
    """
    replacements = {
        "{{MEETING_TITLE}}": mom.meeting_title or "",
        "{{DATE}}": mom.date or "",
        "{{TIME}}": mom.time or "",
        "{{VENUE}}": mom.venue or "",
        "{{CHAIRED_BY}}": mom.chaired_by or "",
    }

    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                _replace_in_paragraph(paragraph, replacements)
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            _replace_in_paragraph(paragraph, replacements)

    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, replacements)


def _replace_in_paragraph(paragraph, replacements: dict):
    if not paragraph.runs:
        return
    combined = "".join(run.text for run in paragraph.runs)
    if not combined:
        return
    updated = combined
    for token, value in replacements.items():
        updated = updated.replace(token, value)
    if updated == combined:
        return
    paragraph.runs[0].text = updated
    for run in paragraph.runs[1:]:
        run.text = ""
